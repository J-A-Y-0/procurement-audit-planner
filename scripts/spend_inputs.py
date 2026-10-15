"""Read the six public workbook fields while preserving their original rows."""
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pandas as pd

from common import normalise_supplier, parse_spend_date

HEADERS = ['Service', 'Expenditure category', 'Narrative', 'Date', 'Net amount', 'Supplier name']


def read_workbooks(paths):
    if not paths:
        raise ValueError('No spend workbooks supplied')
    records = []
    for supplied in sorted(paths):
        path = Path(supplied)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if 'Data to publish' not in workbook.sheetnames:
                raise ValueError(f'{path.name}: sheet Data to publish missing')
            sheet = workbook['Data to publish']
            header = None
            for sheet_row, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if header is None:
                    names = [str(v).strip() if v is not None else '' for v in values]
                    if names[:6] == HEADERS and not any(names[6:]):
                        header = HEADERS
                    elif sheet_row >= 20:
                        raise ValueError(f'{path.name}: required six-column header not found')
                    continue
                if all(v is None for v in values):
                    continue
                row = dict(zip(header, values))
                if row.get('Supplier name') is None or not str(row['Supplier name']).strip():
                    raise ValueError(f'{path.name}:{sheet_row}: supplier missing')
                try:
                    amount = Decimal(str(row['Net amount']))
                except InvalidOperation as error:
                    raise ValueError(f'{path.name}:{sheet_row}: invalid amount') from error
                if not amount.is_finite() or amount < 0:
                    raise ValueError(f'{path.name}:{sheet_row}: amount is negative or non-finite')
                stamp = parse_spend_date(row['Date'])
                row.update({'Net amount': float(amount), 'source_file': path.name,
                            'sheet_row': sheet_row, 'row_id': f'{path.name}:{sheet_row}',
                            'Date': stamp.isoformat(timespec='milliseconds'),
                            'pay_date': stamp.date().isoformat(),
                            'supplier_norm': normalise_supplier(row['Supplier name'])})
                records.append(row)
            if header is None:
                raise ValueError(f'{path.name}: required six-column header not found')
        finally:
            workbook.close()
    return pd.DataFrame(records)
